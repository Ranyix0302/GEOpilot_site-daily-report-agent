from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "网站使用说明书_中文版.docx"

NAVY = "17324D"
BLUE = "2563A5"
LIGHT_BLUE = "EAF3FA"
PALE_BLUE = "F5F9FC"
MID_GRAY = "667085"
LIGHT_GRAY = "D9D9D9"
VERY_LIGHT = "F7F8FA"
BLACK = "000000"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=120, start=150, bottom=120, end=150):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LIGHT_GRAY, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name="Microsoft YaHei", size=None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = instruction
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_end)
    set_run_font(run, size=8.5, color=MID_GRAY)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_label(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text.upper())
    set_run_font(r, size=9, bold=True, color=BLUE)
    return p


def add_title(doc, text, subtitle=None):
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run(text)
    set_run_font(r, size=30, bold=True, color=BLACK)
    if subtitle:
        sp = doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(16)
        sr = sp.add_run(subtitle)
        set_run_font(sr, size=13, color=MID_GRAY)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(10 if level == 1 else 7)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, size=18 if level == 1 else 13, bold=True, color=BLACK)
    return p


def add_body(doc, text, bold_lead=None, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.28
    p.paragraph_format.space_after = Pt(space_after)
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, size=10.5, bold=True, color=BLACK)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, size=10.5, color=BLACK)
    else:
        r = p.add_run(text)
        set_run_font(r, size=10.5, color=BLACK)
    return p


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.left_indent = Inches(0.24 + 0.22 * level)
        p.paragraph_format.first_line_indent = Inches(-0.14)
        p.paragraph_format.line_spacing = 1.18
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        set_run_font(r, size=10.2, color=BLACK)


def add_numbered_steps(doc, steps, start=1):
    for index, (title, detail) in enumerate(steps, start=start):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.34)
        p.paragraph_format.first_line_indent = Inches(-0.34)
        p.paragraph_format.line_spacing = 1.22
        p.paragraph_format.space_after = Pt(7)
        n = p.add_run(f"{index}. ")
        set_run_font(n, size=11, bold=True, color=BLUE)
        t = p.add_run(title)
        set_run_font(t, size=10.5, bold=True, color=BLACK)
        d = p.add_run(f"  {detail}")
        set_run_font(d, size=10.5, color=BLACK)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, value in enumerate(headers):
        cell = header.cells[i]
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(value)
        set_run_font(r, size=9.2, bold=True, color=WHITE)
        if widths:
            cell.width = Inches(widths[i])
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for col_index, value in enumerate(values):
            cell = cells[col_index]
            set_cell_shading(cell, PALE_BLUE if row_index % 2 else WHITE)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            if col_index == 0 and len(value) < 18:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(value)
            set_run_font(r, size=9.2, color=BLACK)
            if widths:
                cell.width = Inches(widths[col_index])
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(1)
    return table


def add_section_intro(doc, number, title, description):
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(6)
    spacer_run = spacer.add_run("\u00A0")
    set_run_font(spacer_run, size=9, color=WHITE)
    add_heading(doc, f"{number}  {title}", 1)
    add_body(doc, description, space_after=10)


def page_break(doc):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.28)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.28)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)

    for name, size in (("Title", 30), ("Heading 1", 18), ("Heading 2", 13)):
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        if name == "Title":
            p_pr = style._element.get_or_add_pPr()
            border = p_pr.find(qn("w:pBdr"))
            if border is not None:
                p_pr.remove(border)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run("SITE DAILY REPORT AGENT  使用说明书")
    set_run_font(hr, size=8, bold=True, color=MID_GRAY)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("工程日报处理与审核  |  第 ")
    set_run_font(fr, size=8.5, color=MID_GRAY)
    add_field(fp, "PAGE")
    fr2 = fp.add_run(" 页")
    set_run_font(fr2, size=8.5, color=MID_GRAY)


def build():
    doc = Document()
    configure_document(doc)

    # Cover
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(42)
    add_label(doc, "工程日报网页操作手册")
    add_title(doc, "Site Daily Report Agent\n网站使用说明书", "适用于工程师 日报审核人员和项目管理人员")
    add_body(
        doc,
        "本手册说明如何使用网页处理施工审核群聊、Operation 群聊和 Silo Operation 群聊资料，如何检查识别结果，以及如何将确认后的数据同步到飞书。",
        space_after=18,
    )
    add_table(
        doc,
        ["项目", "说明"],
        [
            ["网页入口", "服务器电脑使用 http://127.0.0.1:8766/；其他人员使用管理员提供的访问地址"],
            ["日报窗口", "所选日期当天 07:00 至次日 07:00"],
            ["输入资料", "施工 TXT、Operation ZIP、Silo Operation ZIP、可选的操作员映射 Excel"],
            ["输出资料", "施工审核 Excel 和 ZIP、Operation 审核 Word、Silo 审核 Word、飞书日报记录"],
        ],
        widths=[1.35, 5.55],
    )
    add_body(doc, "重要原则：所有 AI 识别结果都应由工程师对照审核材料和原图确认后再同步。", bold_lead="重要原则：", space_after=4)
    add_body(doc, "版本说明：本手册按当前中英文双语网页编写，网页首次打开默认显示英文。", bold_lead="版本说明：")

    # Page 2 quick start
    page_break(doc)
    add_section_intro(doc, 1, "五分钟快速上手", "如果服务已经配置好，可按下面顺序完成一份日报。三个群聊可以分开处理，也可以只处理已有资料的部分。")
    add_numbered_steps(
        doc,
        [
            ("打开网页并选择语言", "右上角选择 EN 或 中文。首次打开默认英文，系统会记住最后一次选择。"),
            ("选择日报日期", "确认页面显示的处理窗口为当天 07:00 至次日 07:00。改变日期会清空当前待处理选择。"),
            ("处理施工审核群聊", "只给当天实际工作的 DCM 设备选择 TXT，核对设备与文件对应关系后再确认处理。"),
            ("处理 Operation 和 Silo", "分别上传对应 ZIP。先核对文件名，再点击确认；两组资料可独立处理。"),
            ("审核并修正", "查看预览表、内容预警和下载的审核材料；双击表格单元格可修改。"),
            ("确认和同步", "确认 Operation 阶段后查看操作员分析，最后在综合预览中保存并同步飞书。"),
        ],
    )
    add_heading(doc, "资料与页面入口对照", 2)
    add_table(
        doc,
        ["页面入口", "应上传的资料", "处理结果"],
        [
            ["施工审核群聊", "每台工作设备各 1 个 WhatsApp TXT", "时间和桩号记录、审核 Excel、结构化审核 ZIP"],
            ["Operation 群聊", "含聊天 TXT 和图片的 ZIP", "机器屏幕桩号、操作员、审核 Word"],
            ["Silo Operation 群聊", "含聊天 TXT 和图片的 ZIP", "水泥仪表读数、除以 2 后的 Cement Content、审核 Word"],
            ["综合预览与写入", "不需重新上传", "合并预览、预警、飞书新增或更新"],
        ],
        widths=[1.55, 2.4, 3.0],
    )

    # Page 3 preparation
    page_break(doc)
    add_section_intro(doc, 2, "开始前的准备", "网页由服务器电脑运行。普通工程师只需准备当天导出的群聊文件，不需要接触 API Key 或飞书密钥。")
    add_heading(doc, "管理员启动服务", 2)
    add_numbered_steps(
        doc,
        [
            ("打开项目文件夹", "进入 agent_v2 文件夹，在空白位置打开 PowerShell。"),
            ("运行启动脚本", "输入 .\\start.ps1 并按 Enter。服务启动后不要关闭这个 PowerShell 窗口。"),
            ("打开网页", "在服务器电脑浏览器访问 http://127.0.0.1:8766/。需要停止服务时回到 PowerShell 按 Ctrl+C。"),
        ],
    )
    add_heading(doc, "工程师准备资料", 2)
    add_bullets(
        doc,
        [
            "确认日报日期，并明确当天实际工作的 DCM 设备编号。",
            "施工审核群聊导出为 TXT；Operation 和 Silo Operation 分别导出为含聊天 TXT 和原始图片的 ZIP。",
            "不要手工删除 ZIP 内的图片或修改聊天 TXT 的时间行，否则可能影响图片与消息匹配。",
            "文件较大时等待上传完成，不要连续重复点击确认按钮。默认单个上传文件上限为 100 MB。",
        ],
    )
    add_heading(doc, "状态检查", 2)
    add_body(doc, "页面右上角会显示模型连接状态。出现 DeepSeek 已就绪后才能进行图片识别；如果显示等待 DEEPSEEK_API_KEY，应联系管理员检查 .env 并重启服务。", bold_lead="页面右上角会显示模型连接状态。")
    add_table(
        doc,
        ["状态", "含义", "处理办法"],
        [
            ["DeepSeek 已就绪", "图片识别服务可用", "继续上传资料"],
            ["等待 API Key", "服务器未读取到 DeepSeek Key", "管理员检查 .env 后重启"],
            ["处理中", "后台正在识别图片", "可切换到其他群聊页面等待"],
            ["处理失败", "文件或模型请求出现错误", "阅读错误提示，检查文件后重试"],
        ],
        widths=[1.4, 2.55, 3.0],
    )

    # Page 4 interface
    page_break(doc)
    add_section_intro(doc, 3, "认识网页界面", "页面由顶部控制区、左侧流程导航和主工作区组成。左侧各阶段相互独立，不要求固定顺序。")
    add_table(
        doc,
        ["区域", "作用", "使用要点"],
        [
            ["右上语言开关", "切换英文和中文", "首次为英文；切换不会清空已选文件和处理结果"],
            ["日报日期", "建立本次 24 小时处理窗口", "必须先选日期；变更日期后需重新选择资料"],
            ["左侧流程导航", "进入施工、Operation、Silo、综合预览和分析", "处理中可以切换页面，后台任务会继续"],
            ["预览表", "显示合并后的八个飞书字段", "双击单元格修改，修改后记得保存"],
            ["内容预警", "显示缺失、冲突和时间异常", "预警不会自动写入飞书，需要人工核对"],
            ["顶部下载按钮", "下载当前阶段审核材料", "阶段处理成功前按钮保持锁定"],
        ],
        widths=[1.65, 2.2, 3.1],
    )
    add_heading(doc, "推荐处理顺序", 2)
    add_body(doc, "建议先处理施工 TXT，再处理 Operation 和 Silo。这样系统能更早发现图片桩号与施工记录不匹配的问题。但如果资料尚未全部收到，也可以先处理任何一个群聊，并先同步已有字段。")
    add_heading(doc, "预览表的八个字段", 2)
    add_table(
        doc,
        ["字段", "来源", "字段", "来源"],
        [
            ["Date", "日报日期", "Operator", "Operation 账号与映射表"],
            ["Pile NO.", "施工 TXT 或群聊资料", "Cement Content", "Silo LED 读数除以 2"],
            ["Equipment NO.", "所选 DCM 设备卡片", "Drilling Start Time", "施工 TXT"],
            ["Drilling Complete Time", "施工 TXT", "Point Complete Time", "施工 TXT"],
        ],
        widths=[1.45, 2.05, 1.55, 1.85],
    )

    # Page 5 construction
    page_break(doc)
    add_section_intro(doc, 4, "处理施工审核群聊", "施工阶段按 DCM-1 至 DCM-8 分配 TXT。只选择当天实际工作的设备；未选择设备会被视为当天未工作，不会产生缺失预警。")
    add_numbered_steps(
        doc,
        [
            ("进入施工审核群聊", "确认页面显示的日报日期正确。"),
            ("选择工作设备文件", "在对应 DCM 卡片中选择 TXT。文件属于哪台设备，就必须放进哪张设备卡片。"),
            ("检查文件清单", "页面会显示已选设备数量和文件名。发现选错时先取消选择，再重新添加。"),
            ("确认并处理", "点击确认文件并开始处理。系统一次性上传所选 TXT，并提取桩号和三个施工时间。"),
            ("检查结果", "查看预览表和内容预警，重点核对缺失时间、时间顺序异常和重复桩号。"),
            ("下载审核材料", "下载施工审核 Excel 和结构化审核 ZIP，留作人工复核或归档。"),
        ],
    )
    add_heading(doc, "施工阶段检查清单", 2)
    add_bullets(
        doc,
        [
            "Equipment NO. 是否与 DCM 卡片一致。",
            "Pile NO. 是否完整，字母、数字和连字符是否正确。",
            "三个时间是否都在所选日报窗口内，且开始时间早于钻进完成和桩点完成时间。",
            "同一桩号是否出现不同设备或不同时间内容。",
        ],
    )
    add_body(doc, "如果当天只有部分设备工作，不要为未工作的设备上传空文件。系统会把未选择设备明确标记为当天未工作。", bold_lead="如果当天只有部分设备工作，")

    # Page 6 operation
    page_break(doc)
    add_section_intro(doc, 5, "处理 Operation 群聊", "Operation 阶段从机器控制屏图片读取完整桩号，并根据 WhatsApp 消息发送账号匹配操作员全名。")
    add_numbered_steps(
        doc,
        [
            ("准备 ZIP", "ZIP 内必须同时有 WhatsApp 导出的 TXT 和对应原始图片。"),
            ("选择并核对", "进入 Operation 群聊，拖入或选择 ZIP。先确认文件名和文件大小是否正确。"),
            ("开始识别", "点击确认文件并开始处理。上传后可切换到其他页面，任务会在后台继续。"),
            ("打开审核 Word", "处理完成后下载 Operation 审核 Word，对照发送账号、全名、发送时间、原图和识别桩号。"),
            ("修正预览", "如桩号或操作员不正确，双击预览表对应单元格修改并保存。"),
            ("确认本阶段", "确认后操作员工作分析会更新；未确认的数据不应作为最终分析依据。"),
        ],
    )
    add_heading(doc, "操作员映射表", 2)
    add_body(doc, "若页面出现操作员全名未匹配，可通过操作员映射表入口上传 Excel。Excel 表头必须包含 Account Name 和 Full Name；Account Name 应与聊天记录中的发送账号一致。")
    add_heading(doc, "识别结果需要重点复核的情况", 2)
    add_table(
        doc,
        ["情况", "应采取的动作"],
        [
            ["低置信度", "打开审核 Word 放大原图，逐字符核对桩号"],
            ["同一桩号多张图片", "判断是重复上传还是确有多条证据，避免重复记录"],
            ["桩号未匹配施工记录", "检查图片桩号和施工 TXT；资料不全时可保留预警等待补充"],
            ["操作员未匹配", "检查账号拼写或更新映射 Excel，再重新处理"],
        ],
        widths=[1.75, 5.2],
    )

    # Page 7 silo
    page_break(doc)
    add_section_intro(doc, 6, "处理 Silo Operation 群聊", "Silo 阶段从聊天消息读取桩号，从水泥计数器图片读取真正点亮的红色 LED 数字，并把读数除以 2 写入 Cement Content。")
    add_numbered_steps(
        doc,
        [
            ("准备 ZIP", "保留原始聊天 TXT 和原始仪表图片，不要先手工裁剪图片。"),
            ("选择并确认", "进入 Silo Operation 群聊，选择正确 ZIP，核对文件名后开始处理。"),
            ("等待识别", "系统自动定位红色 LED 区域；无法得到有效结果时会回退到原图。"),
            ("审核原图", "下载 Silo 审核 Word，核对消息桩号、原始仪表图片、识别读数和除以 2 后的结果。"),
            ("修改错误", "若预览中的 Cement Content 不正确，双击单元格改为人工确认值并保存。"),
        ],
    )
    add_heading(doc, "红色 LED 数字的核对方法", 2)
    add_bullets(
        doc,
        [
            "只读取真正明亮点亮的数字，不要把未点亮但仍可见的暗红色 8 字轮廓计入读数。",
            "例如前两位只是暗淡轮廓、后两位清晰显示 33，则原始读数是 33，不是 8833 或 8888。",
            "机械水表、红色指针、时间戳和设备标签不是目标读数。",
            "确认原始读数后，再检查 Cement Content 是否等于该读数除以 2。",
        ],
    )
    add_table(
        doc,
        ["原始 LED 读数", "网页 Cement Content"],
        [["10", "5"], ["33", "16.5"], ["6", "3"], ["21", "10.5"]],
        widths=[3.45, 3.45],
    )

    # Page 8 review and sync
    page_break(doc)
    add_section_intro(doc, 7, "审核预览并同步飞书", "综合预览会按 Pile NO. 合并当前已处理的资料。同步前应完成表格、预警和审核材料三项检查。")
    add_heading(doc, "预览与修改", 2)
    add_numbered_steps(
        doc,
        [
            ("进入综合预览与写入", "确认当前已有需要同步的记录。"),
            ("逐行检查", "重点检查桩号、设备、时间、操作员和水泥用量是否来自同一桩。"),
            ("处理预警", "预警不会自动阻止同步，也不会自动写入飞书；必须由工程师判断并修正。"),
            ("保存修改", "双击单元格修改后，点击保存当前修改。"),
            ("同步飞书", "点击保存预览并同步飞书，阅读确认提示后继续。"),
        ],
    )
    add_heading(doc, "飞书写入规则", 2)
    add_bullets(
        doc,
        [
            "系统以 Pile NO. 作为唯一主键。飞书没有该桩号时新增，有该桩号时更新本次已有值。",
            "只处理一类或两类群聊也可以同步；后续资料会继续更新同一桩号，不会主动新增重复行。",
            "如果飞书中同一桩号已经存在多行，系统会停止写入，需要人工先合并重复记录。",
            "如新操作员不在飞书 Operator 单选项中，网页会先请求确认；确认后再添加选项并继续。",
        ],
    )
    add_heading(doc, "同步前最终检查", 2)
    add_table(
        doc,
        ["检查项", "通过标准"],
        [
            ["日报日期", "日期与 07:00 至次日 07:00 窗口正确"],
            ["桩号", "格式完整，无明显重复或错别字符"],
            ["施工时间", "字段齐全，先后顺序合理"],
            ["操作员", "全名正确，账号映射已核对"],
            ["水泥用量", "与 Silo 原始 LED 读数除以 2 一致"],
            ["预警", "每条都已查看并作出处理决定"],
        ],
        widths=[1.8, 5.1],
    )

    # Page 9 troubleshooting/security
    page_break(doc)
    add_section_intro(doc, 8, "常见问题和安全要求", "遇到问题时，先保留原始文件和错误提示，再按下表处理。不要为了消除错误而修改原始证据。")
    add_table(
        doc,
        ["问题", "可能原因", "解决办法"],
        [
            ["网页打不开", "服务未启动或 PowerShell 已关闭", "管理员重新运行 .\\start.ps1，并保持窗口开启"],
            ["图片识别不可用", "DeepSeek Key 缺失或服务未重启", "管理员检查 .env，保存后停止并重启服务"],
            ["ZIP 无法处理", "缺少聊天 TXT、ZIP 损坏或日期窗口无消息", "重新从群聊导出并确认日报日期"],
            ["上传后一直处理中", "图片较多或网络请求较慢", "等待并查看进度；不要重复上传同一文件"],
            ["识别结果不准确", "图片模糊、目标区域过小或暗色残影干扰", "对照审核 Word 原图人工修正；保留低置信度预警"],
            ["飞书同步失败", "凭证、权限、字段名或字段类型不正确", "联系管理员检查飞书应用和八个字段配置"],
            ["重启后记录消失", "当前会话数据保存在服务器内存", "重新选择日报日期并上传原始资料"],
        ],
        widths=[1.55, 2.35, 3.0],
    )
    add_heading(doc, "数据与凭证安全", 2)
    add_bullets(
        doc,
        [
            "不要把 .env、DeepSeek API Key、飞书 App Secret 发给工程师或放入前端代码。",
            "工程师只通过网页上传资料；浏览器不会要求输入 API Key。",
            "上传图片和审核材料会保存在服务器本机 data 文件夹。按项目的数据管理要求定期归档或清理。",
            "网页会话在服务重启后清空，因此同步或下载前不要随意关闭服务。",
            "AI 识别用于减少人工录入，不代替工程师对原始资料的最终审核责任。",
        ],
    )
    add_heading(doc, "需要联系管理员的情况", 2)
    add_body(doc, "网页无法启动、模型长期未就绪、飞书无法连接、字段配置不一致、文件超过上传上限，或同一桩号在飞书中存在多条重复记录时，应由管理员处理。联系时请提供日报日期、阶段名称、文件名和完整错误提示，但不要发送密钥。")

    doc.core_properties.title = "Site Daily Report Agent 网站使用说明书"
    doc.core_properties.subject = "工程日报网页中文操作手册"
    doc.core_properties.author = "Site Daily Report Agent"
    doc.core_properties.keywords = "工程日报, 使用说明, Operation, Silo, 飞书"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
